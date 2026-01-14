import streamlit as st
import pandas as pd
import numpy as np
import re
import os
import docx
import PyPDF2
import warnings
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import LabelEncoder
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.calibration import CalibratedClassifierCV
from sklearn.metrics import accuracy_score
from fpdf import FPDF

warnings.filterwarnings("ignore")

DATASET_PATH = "Dataset/resume_dataset.csv"
EVAL_FOLDER = "Model_Evaluation"

def export_as_pdf(name, contact, summary, skills, experience, education):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    pdf.set_font("Times", "B", 24)
    pdf.cell(0, 10, name.upper(), ln=True, align="C")
    pdf.set_font("Times", "", 10)
    pdf.cell(0, 5, contact, ln=True, align="C")
    pdf.ln(5)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)

    sections = [
        ("PROFESSIONAL SUMMARY", summary),
        ("TECHNICAL EXPERTISE", skills),
        ("PROFESSIONAL EXPERIENCE", experience),
        ("EDUCATION", education)
    ]

    for title, content in sections:
        pdf.set_font("Times", "B", 12)
        pdf.set_fill_color(240, 240, 240)
        pdf.cell(0, 8, title, ln=True, fill=True)
        pdf.ln(2)
        pdf.set_font("Times", "", 11)
        pdf.multi_cell(0, 6, content)
        pdf.ln(5)
        
    return pdf.output(dest="S").encode("latin-1")

def extract_text(file):
    extension = file.name.split(".")[-1].lower()
    if extension == "pdf":
        reader = PyPDF2.PdfReader(file)
        text = ""
        for page in reader.pages:
            if page.extract_text():
                text += page.extract_text()
        return text
    elif extension == "docx":
        doc = docx.Document(file)
        return " ".join(p.text for p in doc.paragraphs)
    elif extension == "txt":
        return file.read().decode("utf-8", errors="ignore")
    return ""

def clean_text(text):
    text = text.lower()
    text = re.sub(r"http\S+", " ", text)
    text = re.sub(r"[^a-z\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def get_chatbot_recommendations(text, role):
    recommendations = []
    if not re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text):
        recommendations.append("Contact Information: Missing or unreadable email address.")
    
    word_count = len(text.split())
    if word_count < 200:
        recommendations.append("Resume Length: Content is brief. Consider adding more detail.")
    
    skill_suggestions = {
        "Data Science": ["Python", "Machine Learning", "Pandas", "SQL", "Statistics"],
        "Java Developer": ["Spring Boot", "Hibernate", "Microservices", "Java"],
        "Web Designing": ["HTML5", "CSS3", "JavaScript", "React"],
        "HR": ["Recruitment", "Employee Engagement", "Payroll"],
        "Testing": ["Selenium", "Automation", "Junit"]
    }
    
    suggested_skills = skill_suggestions.get(role, ["Certifications", "Soft Skills"])
    missing = [skill for skill in suggested_skills if skill.lower() not in text.lower()]
    if missing:
        recommendations.append(f"Skill Gap: For {role}, consider adding: {', '.join(missing)}.")
    
    return recommendations

@st.cache_resource
def train_model():
    if not os.path.exists(DATASET_PATH): return None, None, None
    data = pd.read_csv(DATASET_PATH)
    data["cleaned"] = data["Resume"].apply(clean_text)
    
    encoder = LabelEncoder()
    y = encoder.fit_transform(data["Category"])
    
    X_train, X_test, y_train, y_test = train_test_split(
        data["cleaned"], y, test_size=0.2, random_state=42, stratify=y
    )
    
    vectorizer = TfidfVectorizer(stop_words="english", max_features=15000, ngram_range=(1,2))
    X_train_vec = vectorizer.fit_transform(X_train)
    X_test_vec = vectorizer.transform(X_test)
    
    model = CalibratedClassifierCV(LogisticRegression(max_iter=1000, class_weight="balanced"), cv=5)
    model.fit(X_train_vec, y_train)
    
    return model, vectorizer, encoder

def main():
    st.set_page_config(page_title="JobExpert Pro", layout="wide")
    st.markdown("<style>.stApp {background-color: #0d1117; color: #e6e6e6;}</style>", unsafe_allow_html=True)
    
    model, vectorizer, encoder = train_model()
    if not model:
        st.error(f"Dataset not found at '{DATASET_PATH}'")
        return

    st.sidebar.title("JobExpert Pro")
    page = st.sidebar.radio("Services", ["Resume Scanner", "CV Maker"])

    if page == "Resume Scanner":
        st.title("JobExpert Pro: Scanner")
        st.write("AI-powered resume screening and job role classification.")
        
        t1, t2 = st.tabs(["Upload Resume", "Paste Text"])
        input_text = ""
        with t1:
            f = st.file_uploader("Upload PDF or DOCX", type=["pdf", "docx", "txt"])
            if f: input_text = extract_text(f)
        with t2:
            p = st.text_area("Paste Content Here", height=250)
            if p: input_text = p

        if st.button("Analyze Resume"):
            if input_text:
                cleaned = clean_text(input_text)
                vec = vectorizer.transform([cleaned])
                probs = model.predict_proba(vec)[0]
                idx = np.argmax(probs)
                role, conf = encoder.classes_[idx], probs[idx]*100

                col_res, col_bot = st.columns(2)
                with col_res:
                    st.subheader("Screening Result")
                    st.metric("Role", role)
                    #st.metric("Confidence", f"{conf:.2f}%")
                    #st.progress(conf/100)
                
                with col_bot:
                    st.subheader("AI Advisor Feedback")
                    recs = get_chatbot_recommendations(input_text, role)
                    for r in recs: st.write(f"- {r}")
            else:
                st.warning("Please provide resume text.")

    else:
        st.title("CV Maker")
        st.write("Build a ATS-optimized executive resume.")
        
        col_in, col_pre = st.columns([1, 1.2], gap="large")
        
        with col_in:
            with st.form("builder_form"):
                st.subheader("Personal Details")
                name = st.text_input("Full Name", "Husnain Rehman")
                contact = st.text_input("Contact Details", "Lahore, Pakistan | husnainrehman26@gmail.com")
                target_role = st.selectbox("Target Job Category", encoder.classes_)
                summary = st.text_area("Professional Summary", height=100)
                skills = st.text_area("Technical Expertise", height=80)
                experience = st.text_area("Professional Experience", height=150)
                education = st.text_area("Education", height=80)
                submit = st.form_submit_button("Generate Professional CV")

        with col_pre:
            if submit:
                content = f"{target_role} {summary} {skills} {experience}"
                vec = vectorizer.transform([clean_text(content)])
                score = model.predict_proba(vec)[0][list(encoder.classes_).index(target_role)] * 100
                
                if score >= 90:
                    st.success(f"ATS Score: {score:.1f}% - Optimized!")
                else:
                    st.warning(f"ATS Optimization Score: {score:.1f}%. Add keywords related to {target_role} to reach 90%.")

                cv_template = f"""
                <div style="background-color: white; color: #1b1f23; padding: 40px; font-family: 'Times New Roman', serif; border-radius: 4px; line-height: 1.5; border: 1px solid #ddd;">
                    <div style="text-align: center; border-bottom: 2px solid #222; margin-bottom: 15px; padding-bottom: 10px;">
                        <h1 style="margin: 0; font-size: 26px; color: #000; text-transform: uppercase;">{name}</h1>
                        <p style="margin: 5px 0; font-size: 13px; color: #444;">{contact}</p>
                    </div>
                    <h3 style="font-size: 16px; border-bottom: 1px solid #333; text-transform: uppercase; margin-top: 20px;">Professional Summary</h3>
                    <p style="font-size: 14px; white-space: pre-line;">{summary}</p>
                    <h3 style="font-size: 16px; border-bottom: 1px solid #333; text-transform: uppercase; margin-top: 20px;">Technical Expertise</h3>
                    <p style="font-size: 14px; white-space: pre-line;">{skills}</p>
                    <h3 style="font-size: 16px; border-bottom: 1px solid #333; text-transform: uppercase; margin-top: 20px;">Experience</h3>
                    <p style="font-size: 14px; white-space: pre-line;">{experience}</p>
                    <h3 style="font-size: 16px; border-bottom: 1px solid #333; text-transform: uppercase; margin-top: 20px;">Education</h3>
                    <p style="font-size: 14px; white-space: pre-line;">{education}</p>
                </div>
                """
                st.markdown(cv_template, unsafe_allow_html=True)
                
                pdf_data = export_as_pdf(name, contact, summary, skills, experience, education)
                st.download_button(
                    label="Download CV as PDF",
                    data=pdf_data,
                    file_name=f"{name.replace(' ', '_')}_CV.pdf",
                    mime="application/pdf"
                )
            else:
                st.info("Fill the form and click generate to preview and download your CV.")

if __name__ == "__main__":
    main()